import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Asset paths
IMG_ADVISOR = r"C:\Users\sesef\.gemini\antigravity-ide\brain\509906f2-9a3c-418d-bd1e-f90eef3f5416\field_advisor_gps_1789670442100.jpg"
IMG_CHECKIN = r"C:\Users\sesef\.gemini\antigravity-ide\brain\509906f2-9a3c-418d-bd1e-f90eef3f5416\mobile_visit_checkin_1789670457075.jpg"
IMG_OFFLINE = r"C:\Users\sesef\.gemini\antigravity-ide\brain\509906f2-9a3c-418d-bd1e-f90eef3f5416\offline_sync_tech_1789670473483.jpg"

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # AgricoVet Color Palette
    PRIMARY_DARK = RGBColor(15, 23, 42)      # Slate 900
    TEAL_DARK = RGBColor(15, 118, 110)       # Teal 700
    TEAL_PRIMARY = RGBColor(13, 148, 136)    # Teal 600
    TEAL_LIGHT = RGBColor(204, 251, 241)     # Teal 100
    EMERALD = RGBColor(16, 185, 129)         # Emerald 500
    EMERALD_LIGHT = RGBColor(236, 253, 245)  # Emerald 50
    CARD_BG = RGBColor(248, 250, 252)        # Slate 50
    CARD_BORDER = RGBColor(226, 232, 240)    # Slate 200
    TEXT_MAIN = RGBColor(30, 41, 59)         # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139)     # Slate 500
    WHITE = RGBColor(255, 255, 255)
    AMBER = RGBColor(245, 158, 11)           # Amber 500
    AMBER_LIGHT = RGBColor(254, 243, 199)    # Amber 100
    BLUE = RGBColor(59, 130, 246)            # Blue 500
    BLUE_LIGHT = RGBColor(239, 246, 255)     # Blue 50
    PURPLE = RGBColor(139, 92, 246)          # Purple 500
    PURPLE_LIGHT = RGBColor(245, 243, 255)   # Purple 50

    def add_header(slide, title_text, subtitle_text, category="MÓDULO DE GESTIÓN DE VISITAS"):
        # Header accent band
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
        band.fill.solid()
        band.fill.fore_color.rgb = TEAL_PRIMARY
        band.line.fill.background()

        # Category pill background
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.4), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = TEAL_LIGHT
        pill.line.color.rgb = TEAL_PRIMARY
        pill.line.width = Pt(1)
        tf_pill = pill.text_frame
        tf_pill.word_wrap = False
        tf_pill.margin_left = Inches(0.15)
        tf_pill.margin_top = Inches(0.02)
        p_pill = tf_pill.paragraphs[0]
        p_pill.text = f"AGRICOMET  •  {category.upper()}"
        p_pill.font.size = Pt(9.5)
        p_pill.font.bold = True
        p_pill.font.color.rgb = TEAL_DARK

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.65))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(25)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.36), Inches(11.7), Inches(0.45))
        tf = sub_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED

    def add_card(slide, left, top, width, height, title, items, badge_text=None, border_color=CARD_BORDER, bg_color=WHITE, accent_bar_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        if accent_bar_color:
            top_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, Inches(0.12))
            top_bar.fill.solid()
            top_bar.fill.fore_color.rgb = accent_bar_color
            top_bar.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.28)
        tf.margin_bottom = Inches(0.25)

        if badge_text:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge_text.upper()
            p_badge.font.size = Pt(9)
            p_badge.font.bold = True
            p_badge.font.color.rgb = accent_bar_color if accent_bar_color else TEAL_PRIMARY
            p_badge.space_after = Pt(4)
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(15)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY_DARK
        p_title.space_after = Pt(8)

        for item in items:
            p_item = tf.add_paragraph()
            p_item.text = f"•  {item}"
            p_item.font.size = Pt(11)
            p_item.font.color.rgb = TEXT_MAIN
            p_item.space_after = Pt(5)

        return card

    # ==========================================
    # SLIDE 1: PORTADA IMPACTANTE CON ILUSTRACIÓN 3D
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY_DARK
    bg1.line.fill.background()

    # Left text column
    tb = slide1.shapes.add_textbox(Inches(0.9), Inches(1.3), Inches(6.2), Inches(5.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "AGRICOMET DE GUATEMALA"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.space_after = Pt(14)

    p1 = tf.add_paragraph()
    p1.text = "Módulo de Rutas y Visitas a Clientes"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "Capacitación Operativa y Guía de Campo para Asesores Técnicos"
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_after = Pt(28)

    p3 = tf.add_paragraph()
    p3.text = "📍 Geolocalización Satelital  •  📸 Evidencias Fotográficas  •  🗺️ Google Earth en Vivo"
    p3.font.size = Pt(11.5)
    p3.font.bold = True
    p3.font.color.rgb = TEAL_LIGHT

    # Right side: 3D Illustration
    if os.path.exists(IMG_ADVISOR):
        slide1.shapes.add_picture(IMG_ADVISOR, Inches(7.0), Inches(1.0), width=Inches(5.7))

    # ==========================================
    # SLIDE 2: EL PROPÓSITO & MÉTRICAS
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "¿Por qué este módulo? Beneficios Clave para el Asesor", "Dejamos atrás las libretas y mensajes dispersos de WhatsApp para potenciar tu trabajo en ruta.")

    # 3 Stat Badges at top
    stats_data = [
        ("⏱️ -40% TIEMPO", "Rutas optimizadas por cercanía geográfica sin dar vueltas"),
        ("🛡️ 100% RESPALDO", "Prueba irrefutable con foto, GPS y hora exacta de cada visita"),
        ("💰 +25% CIERRES", "Mayor frecuencia de contacto y cobranza al día")
    ]
    for i, (metric, label) in enumerate(stats_data):
        box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + i * 4.0), Inches(1.9), Inches(3.7), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = TEAL_LIGHT
        box.line.color.rgb = TEAL_PRIMARY
        box.line.width = Pt(1.5)
        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0.18)
        tf_b.margin_top = Inches(0.1)
        p_m = tf_b.paragraphs[0]
        p_m.text = metric
        p_m.font.size = Pt(15)
        p_m.font.bold = True
        p_m.font.color.rgb = TEAL_DARK
        p_l = tf_b.add_paragraph()
        p_l.text = label
        p_l.font.size = Pt(10)
        p_l.font.color.rgb = TEXT_MAIN

    # 3 Content Cards below
    cards_data = [
        ("Organización de Ruta", [
            "Visualiza tus clientes en el mapa según cercanía en km.",
            "Ahorra tiempo y combustible en cada viaje.",
            "Filtra clientes pendientes de visita este mes."
        ], "EFICIENCIA EN CAMPO", TEAL_PRIMARY),
        ("Respaldo Total de tu Trabajo", [
            "Cada parada queda sellada con foto y coordenadas GPS.",
            "Prueba oficial de cobranza, asesoría y entrega.",
            "Reconocimiento transparente de tu esfuerzo diario."
        ], "TRANSPARENCIA Y SEGURIDAD", BLUE),
        ("Cierre de Ventas y Cobranza", [
            "Conoce el saldo y compras antes de bajarte del vehículo.",
            "Programa tu próxima visita para no descuidar clientes.",
            "Directamente sincronizado con ventas y recibos de caja."
        ], "MÁS VENTAS", EMERALD)
    ]
    for i, (title, items, badge, accent_color) in enumerate(cards_data):
        add_card(slide2, Inches(0.8 + i * 4.0), Inches(3.05), Inches(3.7), Inches(3.9), title, items, badge, accent_bar_color=accent_color)

    # ==========================================
    # SLIDE 3: PASO CERO: ACTIVAR GPS
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "Paso Cero: Activar el GPS en tu Teléfono", "La aplicación necesita tu ubicación precisa para marcar clientes y certificar visitas.")

    add_card(slide3, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), "📱 En Teléfonos Android", [
        "Desliza hacia abajo la barra superior y activa 'Ubicación' (GPS).",
        "Abre la app de AgricoVet o ingresa a https://agricovet.lat.",
        "Cuando el navegador pregunte: '¿Permitir ubicación?', elige 'Permitir mientras se usa la app'.",
        "Ajustes > Apps > Chrome/AgricoVet > Permisos: Activa 'Ubicación precisa'.",
        "✅ Si ves tu punto azul en el mapa, tu GPS está listo y sincronizado."
    ], "ANDROID PASO A PASO", accent_bar_color=EMERALD)

    add_card(slide3, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), "🍎 En Teléfonos iPhone (iOS)", [
        "Ve a Configuración > Privacidad y Seguridad > Localización.",
        "Asegúrate de que 'Localización' esté ENCENDIDA (Verde).",
        "Busca Safari (o Chrome) en la lista de aplicaciones.",
        "Selecciona 'Al usar la app' y enciende el interruptor 'Ubicación precisa'.",
        "Al entrar a la app, toca 'Permitir al usar el sitio web'.",
        "💡 Consejo: Evita usar navegación privada para mantener tu sesión."
    ], "IPHONE / IOS PASO A PASO", accent_bar_color=BLUE)

    # ==========================================
    # SLIDE 4: PANTALLA PRINCIPAL DEL MAPA
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "Conociendo el Mapa de Visitas", "Tu centro de comando visual con todos tus clientes y herramientas clave.")

    add_card(slide4, Inches(0.8), Inches(2.0), Inches(3.7), Inches(4.8), "Barra Superior", [
        "🔘 Filtros de Estado: Ver 'Todos', 'Visitados' (verde) o 'Pendientes' (amarillo).",
        "🔎 Buscador Inteligente: Escribe el nombre, negocio o código del cliente.",
        "🚀 Botón 'Volar a cliente': Te centra en el mapa de inmediato.",
        "📍 Regiones de Guatemala: Accesos directos a Petén, Oriente, Costa Sur, etc."
    ], "NAVEGACIÓN", accent_bar_color=TEAL_PRIMARY)

    add_card(slide4, Inches(4.8), Inches(2.0), Inches(3.7), Inches(4.8), "Pines en el Mapa", [
        "🏢 Pin Verde: Cliente con ubicación GPS ya confirmada.",
        "📍 Pin Azul con Halo: Tu posición GPS actual en tiempo real.",
        "⭕ Círculo sombreado: Precisión de tu señal GPS en metros.",
        "👆 Toca cualquier pin: Muestra nombre, teléfono, ruta y saldo al instante."
    ], "ELEMENTOS VISUALES", accent_bar_color=BLUE)

    add_card(slide4, Inches(8.8), Inches(2.0), Inches(3.7), Inches(4.8), "Botones Flotantes de Acción", [
        "📍 Marcar Cliente: Guarda la ubicación del negocio donde estás parado.",
        "➕ Registrar Visita: Llena la bitácora con foto y motivo de la visita.",
        "🔍 Zoom (+ / -): Acerca hasta nivel de techo o aleja para ver la ruta.",
        "🎯 Mi Ubicación: Te re-centra donde tú estás en este instante."
    ], "ACCIONES RÁPIDAS", accent_bar_color=AMBER)

    # ==========================================
    # SLIDE 5: SELECTOR DE CAPAS DE MAPA
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "Capas de Mapa: Adapta la Vista a tu Terreno", "Diferentes vistas satelitales y viales para cualquier rincón de Guatemala.")

    layers = [
        ("🛰️ Satélite Híbrido", "Foto satelital de alta resolución combinada con nombres de municipios, carreteras y aldeas. Vista recomendada para el 90% de tus rutas.", EMERALD),
        ("🌎 Google Earth (Limpio)", "Satélite 100% puro sin textos ni líneas. Ideal para ubicar techos de bodegas, corrales, cultivos y linderos de fincas.", TEAL_PRIMARY),
        ("🗺️ Calles y Rutas", "Mapa vial clásico y ligero. Perfecto cuando vas manejando en zonas urbanas para ver nombres de calles y avenidas.", BLUE),
        ("⛰️ Relieve y Fincas", "Curvas de nivel, cerros, valles y topografía. Esencial para asesores agrícolas en zonas montañosas.", AMBER),
        ("🛰️ Satélite HD (Esri)", "Vista satelital alternativa con zoom auto-escalado para verificar vegetación y caminos rurales.", PURPLE)
    ]
    for i, (l_title, l_desc, accent) in enumerate(layers):
        top_pos = Inches(2.0 + i * 0.96)
        card = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.82))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        bar_sub = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(0.18), Inches(0.82))
        bar_sub.fill.solid()
        bar_sub.fill.fore_color.rgb = accent
        bar_sub.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = f"{l_title}:  "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK
        run = p.add_run()
        run.text = l_desc
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 6: INICIAR RUTA DEL DÍA
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "Paso 1 del Día: Iniciar tu Ruta de Trabajo", "Al salir en la mañana, abre la app y activa tu jornada con un solo toque.")

    add_card(slide6, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), "¿Cómo se inicia?", [
        "En la parte superior de 'Visitas', verás la tarjeta verde 'Ruta del Día'.",
        "Presiona el botón 'Iniciar Ruta del Día'.",
        "El sistema registrará tu hora exacta de salida y ubicación inicial.",
        "Verás un cronómetro en vivo con el tiempo transcurrido y las paradas que vas realizando.",
        "No necesitas mantener la pantalla encendida todo el camino."
    ], "INICIO DE JORNADA", accent_bar_color=TEAL_PRIMARY)

    add_card(slide6, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), "¿Para qué sirve?", [
        "Auditoría y Justificación: Respalda tus horas trabajadas y recorridos.",
        "Organización: Te ayuda a medir cuántas visitas logras por jornada.",
        "Seguridad: En caso de cualquier percance en carretera, la empresa sabe en qué tramo de ruta estabas.",
        "Historial Mensual: Permite ver la constancia y rendimiento de tus rutas."
    ], "BENEFICIOS CLAVE", accent_bar_color=EMERALD)

    # ==========================================
    # SLIDE 7: MARCAR CLIENTE (GEORREFERENCIACIÓN)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "Función 1: Marcar la Ubicación de un Cliente", "Haz esto UNA SOLA VEZ por cliente cuando estés físicamente en su negocio o finca.")

    steps = [
        ("Paso 1: Llega al lugar", "Estaciónate frente al local, agropecuaria, farmacia o finca del cliente."),
        ("Paso 2: Toca 'Marcar Cliente'", "Presiona el botón flotante 'Marcar Cliente' (con icono de pin de mapa)."),
        ("Paso 3: Selecciona al Cliente", "Búscalo en el listado por nombre comercial, nombre del dueño o código."),
        ("Paso 4: Confirmar GPS", "El sistema toma tus coordenadas automáticamente. Verifica que el indicador de precisión esté en verde."),
        ("Paso 5: Guardar Ubicación", "Presiona 'Guardar Ubicación'. ¡Listo! A partir de ese momento, el cliente tendrá su pin permanente en el mapa.")
    ]
    for i, (s_title, s_desc) in enumerate(steps):
        top_pos = Inches(2.0 + i * 0.96)
        card = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.82))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        # Number circle badge
        num_circle = slide7.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), top_pos + Inches(0.16), Inches(0.5), Inches(0.5))
        num_circle.fill.solid()
        num_circle.fill.fore_color.rgb = TEAL_PRIMARY
        num_circle.line.fill.background()
        tf_num = num_circle.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.alignment = PP_ALIGN.CENTER
        p_num.text = str(i + 1)
        p_num.font.size = Pt(14)
        p_num.font.bold = True
        p_num.font.color.rgb = WHITE

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.85)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = f"{s_title}:  "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEAL_DARK
        run = p.add_run()
        run.text = s_desc
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 8: REGISTRAR VISITA CON MOCKUP 3D
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "Función 2: Registrar una Visita en Terreno", "El registro oficial de cada parada que realizas con el cliente.")

    # Left content: 3 step cards
    v_steps = [
        ("1. Cliente y Motivo", "Toca 'Registrar Visita'. Selecciona cliente y elige: Prospección, Seguimiento, Cobro o Entrega."),
        ("2. Foto y Evidencia", "Toma foto de fachada, mostrador, boleta firmada o producto. La app la comprime automáticamente para cuidar tus datos."),
        ("3. Notas y Guardar", "Anota acuerdos, fecha de la próxima visita y toca 'Guardar Visita'. Queda certificada con fecha, hora y GPS.")
    ]
    for i, (title, desc) in enumerate(v_steps):
        top_pos = Inches(2.0 + i * 1.6)
        card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(6.4), Inches(1.4))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        bar_v = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(0.15), Inches(1.4))
        bar_v.fill.solid()
        bar_v.fill.fore_color.rgb = TEAL_PRIMARY if i==0 else (BLUE if i==1 else EMERALD)
        bar_v.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.15)
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY_DARK
        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MAIN

    # Right side: 3D Checkin Mobile Mockup
    if os.path.exists(IMG_CHECKIN):
        slide8.shapes.add_picture(IMG_CHECKIN, Inches(7.5), Inches(1.8), width=Inches(5.1))

    # ==========================================
    # SLIDE 9: LOS 4 TIPOS DE VISITA (CARDS DE COLORES)
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "Los 4 Tipos de Visita: ¿Cuándo usar cada uno?", "Clasificar correctamente tus visitas te ayuda a medir tu efectividad comercial.")

    types_data = [
        ("💼 Prospección", [
            "Para clientes nuevos o potenciales.",
            "Presentación de catálogo y cotización.",
            "Toma de datos fiscales y teléfono.",
            "Evaluación de compra."
        ], "CLIENTES NUEVOS", EMERALD, EMERALD_LIGHT),
        ("🤝 Seguimiento", [
            "Asesoría técnica en campo.",
            "Revisión de stock AgricoVet en tienda.",
            "Resolución de dudas del cliente.",
            "Fidelización comercial."
        ], "POST-VENTA Y ASESORÍA", BLUE, BLUE_LIGHT),
        ("💰 Cobro", [
            "Cobro de facturas a crédito.",
            "Recojo de cheques o transferencias.",
            "Emisión de Recibo Conforme.",
            "Foto de boleta de depósito."
        ], "CARTERA Y SALDOS", AMBER, AMBER_LIGHT),
        ("📦 Entrega", [
            "Entrega de producto en destino.",
            "Verificación de bultos completos.",
            "Firma de guía o carta de entrega.",
            "Garantía de satisfacción."
        ], "LOGÍSTICA", PURPLE, PURPLE_LIGHT)
    ]
    for i, (title, items, badge, accent, bg) in enumerate(types_data):
        add_card(slide9, Inches(0.8 + i * 3.0), Inches(2.0), Inches(2.75), Inches(4.8), title, items, badge, border_color=accent, bg_color=WHITE, accent_bar_color=accent)

    # ==========================================
    # SLIDE 10: MI CARTERA (CERCANÍA GPS)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "Pestaña 'Mi Cartera': Clientes Ordenados por Cercanía", "Tu teléfono te avisa quién está cerca de ti para que aproveches cada viaje al máximo.")

    add_card(slide10, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), "Cálculo de Distancia en Vivo", [
        "La lista calcula automáticamente a cuántos kilómetros estás de cada cliente (ej. 'a 1.4 km').",
        "Los clientes más cercanos aparecen siempre arriba de la lista.",
        "Si te sobra tiempo entre citas, mira la lista y visita a los que tienes a 5 minutos.",
        "Botones 'Waze' y 'Google Maps': Toca un botón y la app te abre la navegación guiada por voz hasta la puerta del cliente."
    ], "GEOPOSICIÓN INTELIGENTE", accent_bar_color=TEAL_PRIMARY)

    add_card(slide10, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), "Semáforo de Frecuencia de Visita", [
        "🔴 URGENTE (> 30 días sin visita): Clientes que corren riesgo de irse con la competencia.",
        "🟡 REGULAR (15 a 30 días): Clientes en ciclo normal de atención comercial.",
        "🟢 AL DÍA (< 15 días): Clientes atendidos recientemente.",
        "⚪ SIN VISITAS: Nuevos clientes que aún no conocen a su asesor.",
        "Filtra con un solo toque los 'Urgentes' para armar tu plan del día."
    ], "SEMÁFORO COMERCIAL", accent_bar_color=AMBER)

    # ==========================================
    # SLIDE 11: MODO OFFLINE CON DIAGRAMA 3D
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "¿Qué pasa si no hay señal de internet en la finca?", "El sistema está diseñado para trabajar en campo abierto sin cobertura móvil.")

    # Left content
    add_card(slide11, Inches(0.8), Inches(2.0), Inches(5.2), Inches(4.8), "Tecnología Satelital Offline", [
        "1. El GPS SÍ Funciona sin Internet: El chip GPS de tu teléfono conecta directo a los satélites en el espacio sin gastar datos ni saldo.",
        "2. Guardado Local Seguro: La visita se guarda en la memoria interna de tu teléfono con icono '📡 Guardada sin conexión'.",
        "3. Auto-Sincronización Silenciosa: En cuanto vuelvas a carretera con señal 3G/4G/Wi-Fi, la app sube todo automáticamente al servidor.",
        "4. Cero Pérdida de Información: Tus datos y fotos no se pierden aunque apagues el teléfono."
    ], "CERO PÉRDIDA DE DATOS", accent_bar_color=EMERALD)

    # Right side: 3D Offline Sync Diagram
    if os.path.exists(IMG_OFFLINE):
        slide11.shapes.add_picture(IMG_OFFLINE, Inches(6.3), Inches(2.0), width=Inches(6.2))

    # ==========================================
    # SLIDE 12: FINALIZAR RUTA DEL DÍA
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "Paso Final del Día: Finalizar tu Ruta", "Cierra tu jornada laboral y genera el resumen consolidado de tu día.")

    add_card(slide12, Inches(0.8), Inches(2.0), Inches(5.7), Inches(4.8), "Cómo cerrar el día", [
        "Al terminar tu última visita o llegar a casa/oficina, abre la app.",
        "Toca el botón 'Finalizar Ruta del Día'.",
        "Escribe un breve comentario de cierre (ej. 'Ruta completa en La Libertad, 6 clientes visitados, 2 pedidos tomados').",
        "Presiona 'Confirmar Cierre de Ruta'.",
        "El sistema cierra el cronómetro y guarda el resumen completo del día."
    ], "CIERRE DE DÍA", accent_bar_color=TEAL_PRIMARY)

    add_card(slide12, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), "Reporte Automático", [
        "Total de paradas realizadas en el día.",
        "Tiempo total de ruta activa.",
        "Clientes visitados vs pendientes.",
        "Monto total cobrado o vendido durante la jornada.",
        "Tu supervisor recibe el resumen sin que tengas que mandar reportes largos por escrito."
    ], "RESULTADOS", accent_bar_color=EMERALD)

    # ==========================================
    # SLIDE 13: REGLAS DE ORO
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "Reglas de Oro para un Asesor Exitoso", "Buenas prácticas diarias para sacar el máximo provecho a la herramienta.")

    rules = [
        ("1. Marca al cliente en su puerta real", "No marques clientes desde tu casa o el hotel. La ubicación debe ser la del negocio real.", EMERALD),
        ("2. Una foto vale más que mil palabras", "La fotografía es tu mejor respaldo ante cualquier reclamo de cobranza o entrega.", BLUE),
        ("3. Registra la visita en el instante", "No dejes las visitas para la noche. Hacerlo en el instante garantiza hora y GPS exactos.", TEAL_PRIMARY),
        ("4. Cuida tu batería", "Lleva siempre tu cargador de carro o batería portátil (powerbank) en tu vehículo.", AMBER),
        ("5. Programa la siguiente cita", "Siempre deja acordada y anotada en la app la fecha de tu próxima visita.", PURPLE)
    ]
    for i, (r_title, r_desc, accent) in enumerate(rules):
        top_pos = Inches(2.0 + i * 0.96)
        card = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.82))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        bar_r = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(0.18), Inches(0.82))
        bar_r.fill.solid()
        bar_r.fill.fore_color.rgb = accent
        bar_r.line.fill.background()

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.35)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = f"{r_title}:  "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK
        run = p.add_run()
        run.text = r_desc
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 14: DINÁMICA PRÁCTICA EN VIVO
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    add_header(slide14, "Taller Práctico en Vivo: Hagámoslo Juntos", "Todos saquen su teléfono en este momento para hacer una prueba real guiada.")

    steps_live = [
        ("Ejercicio 1: Encender GPS y entrar", "Abre https://agricovet.lat en tu navegador o abre la app. Verifica que tu punto azul aparezca en el mapa."),
        ("Ejercicio 2: Explorar las capas", "Abre el selector arriba a la derecha. Cambia a 'Google Earth' y luego a 'Satélite Híbrido'."),
        ("Ejercicio 3: Iniciar la Ruta", "Presiona 'Iniciar Ruta del Día' para que veas cómo empieza a correr tu contador de jornada."),
        ("Ejercicio 4: Registrar una Visita de Prueba", "Toca 'Registrar Visita', busca al cliente de prueba, tómate una foto o foto de la sala y guarda."),
        ("Ejercicio 5: Ver la visita en el mapa", "Busca el pin que acabas de crear, tócalo y confirma que tus datos y foto quedaron grabados.")
    ]
    for i, (e_title, e_desc) in enumerate(steps_live):
        top_pos = Inches(2.0 + i * 0.96)
        card = slide14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(11.7), Inches(0.82))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = TEAL_PRIMARY
        card.line.width = Pt(1.5)

        circle_e = slide14.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), top_pos + Inches(0.16), Inches(0.5), Inches(0.5))
        circle_e.fill.solid()
        circle_e.fill.fore_color.rgb = TEAL_PRIMARY
        circle_e.line.fill.background()
        tf_e = circle_e.text_frame
        p_ce = tf_e.paragraphs[0]
        p_ce.alignment = PP_ALIGN.CENTER
        p_ce.text = str(i + 1)
        p_ce.font.size = Pt(14)
        p_ce.font.bold = True
        p_ce.font.color.rgb = WHITE

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.85)
        tf.margin_top = Inches(0.12)
        p = tf.paragraphs[0]
        p.text = f"{e_title}:  "
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = TEAL_DARK
        run = p.add_run()
        run.text = e_desc
        run.font.size = Pt(11.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 15: CIERRE Y PREGUNTAS
    # ==========================================
    slide15 = prs.slides.add_slide(blank_layout)
    bg15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = PRIMARY_DARK
    bg15.line.fill.background()

    bar15 = slide15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.15))
    bar15.fill.solid()
    bar15.fill.fore_color.rgb = TEAL_PRIMARY
    bar15.line.fill.background()

    tb15 = slide15.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.9), Inches(4.0))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p0 = tf15.paragraphs[0]
    p0.text = "AGRICOMET  •  GESTIÓN DE RUTAS Y VISITAS"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = EMERALD
    p0.space_after = Pt(20)

    p1 = tf15.add_paragraph()
    p1.text = "¡Éxitos en la Ruta y a Romper Metas!"
    p1.font.size = Pt(38)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(16)

    p2 = tf15.add_paragraph()
    p2.text = "Espacio para Dudas, Preguntas y Comentarios"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.space_after = Pt(28)

    p3 = tf15.add_paragraph()
    p3.text = "La tecnología está a tu servicio para que vendas más, cobres mejor y ganes más comisiones. 🚜🌱"
    p3.font.size = Pt(13.5)
    p3.font.color.rgb = TEAL_LIGHT

    output_path = "Presentacion_Modulo_Visitas_AgricoVet.pptx"
    prs.save(output_path)
    print(f"Deluxe Presentation created successfully at: {output_path}")

if __name__ == '__main__':
    create_deck()
